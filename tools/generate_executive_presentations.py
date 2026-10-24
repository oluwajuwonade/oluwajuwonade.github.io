#!/usr/bin/env python3
"""Generate portfolio-ready PowerPoint samples; CI exports the matching PDFs."""
from pathlib import Path
import os, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = Path(os.environ.get("PRESENTATION_OUT", Path(__file__).resolve().parent / "out")).resolve()
OUT.mkdir(parents=True, exist_ok=True)
PAL={"bg":"0B0D10","panel":"12161B","panel2":"171C22","text":"F2F5F7","muted":"A6B0B9","line":"2A323B","mint":"8FE3CF","amber":"EFB65A","red":"FA8F8F","blue":"87B7FF"}
def c(v):
    v=PAL.get(v,v).lstrip("#"); return RGBColor(int(v[:2],16),int(v[2:4],16),int(v[4:6],16))
def box(s,x,y,w,h,fill="panel",stroke=None,rounded=True):
    q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    q.fill.solid(); q.fill.fore_color.rgb=c(fill); q.line.color.rgb=c(stroke or fill)
    if rounded:
        try:q.adjustments[0]=.12
        except Exception:pass
    return q
def ln(s,x1,y1,x2,y2,color="line",width=.8):
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)); q.line.color.rgb=c(color); q.line.width=Pt(width)
def t(s,txt,x,y,w,h,size=12,color="text",bold=False,align=PP_ALIGN.LEFT,italic=False):
    q=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); f=q.text_frame; f.clear(); f.word_wrap=True
    f.margin_left=Inches(.03); f.margin_right=Inches(.03); f.margin_top=Inches(.025); f.margin_bottom=Inches(.02); f.vertical_anchor=MSO_ANCHOR.TOP
    p=f.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=txt; r.font.name="Aptos"; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic; r.font.color.rgb=c(color)
    return q
def new(prs):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=c("bg"); return s
def footer(s,n):
    ln(s,.6,7.06,12.72,7.06); t(s,"Oluwajuwon Adediji  ·  Data & Quantitative Analyst  ·  Portfolio sample",.62,7.13,9.9,.17,7,"muted")
    t(s,f"{n:02d} / 08",11.75,7.12,.86,.18,7.5,"mint",True,PP_ALIGN.RIGHT)
def head(s,k,title,sub,n):
    t(s,k,.62,.32,11,.22,8.5,"mint",True); t(s,title,.62,.68,12,.54,25,"text",True); t(s,sub,.64,1.29,12,.38,10.5,"muted")
    ln(s,.62,1.78,12.72,1.78); footer(s,n)
def card(s,x,y,w,h,title,body,accent="mint"):
    box(s,x,y,w,h,"panel","line",True); box(s,x,y,.045,h,accent,accent,False)
    t(s,title,x+.22,y+.19,w-.42,.36,13,"text",True)
    t(s,body,x+.22,y+.63,w-.42,h-.76,10.1,"muted")
def kpi(s,x,y,w,label,value,desc,accent="mint"):
    box(s,x,y,w,1.03,"panel","line",True); t(s,label.upper(),x+.15,y+.12,w-.3,.18,7.8,"muted",True)
    t(s,value,x+.15,y+.33,w-.3,.36,21,accent,True); t(s,desc,x+.15,y+.76,w-.3,.2,8,"muted")
def deck(title):
    p=Presentation(); p.slide_width=Inches(13.333); p.slide_height=Inches(7.5)
    p.core_properties.title=title; p.core_properties.author="Oluwajuwon Adediji"
    p.core_properties.subject="Executive presentation portfolio sample"; p.core_properties.keywords="business analysis; quantitative research; decision support"
    return p
def cover(p,kicker,title,sub,side_head,metrics,out_name):
    s=new(p); box(s,9.05,0,4.283,7.5,"panel","panel",False); box(s,9.05,0,.06,7.5,"mint","mint",False)
    t(s,kicker,.72,.62,7.7,.25,9.5,"mint",True); t(s,title,.72,1.18,7.9,1.95,32,"text",True)
    t(s,sub,.75,3.35,7.35,.74,15,"muted")
    for i,tag in enumerate(["BUSINESS ANALYSIS","QUANTITATIVE RESEARCH","DECISION SUPPORT"]):
        xx=.75+i*2.18; box(s,xx,4.42,2.03,.36,"panel2","line",True); t(s,tag,xx+.04,4.53,1.95,.15,7,"mint",True,PP_ALIGN.CENTER)
    t(s,"Prepared by Oluwajuwon Adediji\nData & Quantitative Analyst",.75,5.23,5.7,.57,11.5)
    t(s,"PORTFOLIO SAMPLE · NOT CLIENT RESULTS",.75,6.34,5.8,.2,8,"amber",True)
    t(s,side_head,9.45,.75,3.2,.25,8.5,"mint",True)
    t(s,metrics[0][0],9.45,1.2,3.2,.3,10,"muted"); t(s,metrics[0][1],9.45,1.52,3.15,.45,24,metrics[0][2],True)
    ln(s,9.45,2.1,12.65,2.1)
    t(s,metrics[1][0],9.45,2.45,3.2,.3,10,"muted"); t(s,metrics[1][1],9.45,2.77,3.15,.45,24,metrics[1][2],True)
    t(s,metrics[2],9.45,3.55,3.05,1.0,11,"text")
    footer(s,1)
def cards_slide(p,k,title,sub,n,cards):
    s=new(p); head(s,k,title,sub,n)
    count=len(cards); cols=2 if count!=3 else 3; rows=math.ceil(count/cols); gap=.2
    left=.72; top=2.1; total=11.92; cw=(total-(cols-1)*gap)/cols
    ch=(4.42-(rows-1)*.2)/rows
    for i,item in enumerate(cards):
        row=i//cols; col=i%cols; card(s,left+col*(cw+gap),top+row*(ch+.2),cw,ch,item[0],item[1],item[2] if len(item)>2 else "mint")
def retail():
    p=deck("Retail Sales Performance Diagnostic")
    cover(p,"EXECUTIVE BUSINESS BRIEF","Retail Sales\nPerformance Diagnostic","Executive diagnosis, KPI assurance, and a 30-day decision plan","THE MANAGEMENT QUESTION",
          [("Reported revenue change","-14.8%","red"),("Unit volume","+48.3%","mint"),"Why did revenue and profit fall while volume rose?"])
    cards_slide(p,"01 / EXECUTIVE SUMMARY","Executive readout","Volume improved, but value realization and profit weakened. Treat the driver story as provisional until KPIs reconcile.",2,[
      ("Revenue change","-14.8% reported H1→H2 movement. Confirm period, net/gross definition, returns and filters.","red"),
      ("Unit volume","+48.3% reported H1→H2 movement. Increased volume did not translate into stronger value realization.","mint"),
      ("Average selling price","-33.1% reported movement. Definition and weighting need validation before attributing drivers.","amber"),
      ("Management action","Reconcile revenue, units and ASP first; then investigate dimensions and test interventions with profit-aware guardrails.","blue")])
    s=new(p); head(s,"02 / PERFORMANCE SIGNAL","What the reported KPIs say","H1 = 100 index. The chart is an executive screen for where to investigate, not causal evidence.",3)
    t(s,"H2 INDEX (H1 = 100)",.86,2.05,2.6,.2,8,"mint",True)
    vals=[("Revenue",85.2,"red"),("Units sold",148.3,"mint"),("Average selling price",66.9,"amber"),("Profit",76.6,"red")]
    for i,(lab,val,col) in enumerate(vals):
        yy=2.55+i*.76; t(s,lab,.9,yy,2.15,.25,10.5,"text",True); box(s,3.05,yy,6.2,.32,"panel2","panel2",True)
        box(s,3.05,yy,min(6.2*val/160,6.2),.32,col,col,True); t(s,f"{val:.1f}",9.4,yy-.015,.68,.3,10.5,"text",True)
    card(s,10.08,2.32,2.42,3.18,"Signal","More units did not translate into stronger reported revenue or profit. Check price/discount pressure, mix shifts, and inconsistent definitions.","amber")
    t(s,"Source figures describe a synthetic portfolio case. Validate before using the narrative externally.",.87,6.1,11.6,.28,9,"muted",False,italic=True)
    cards_slide(p,"03 / DATA QUALITY","KPI assurance gate: reconcile before explaining","Under Revenue = Units × weighted ASP, the published headline changes do not reconcile.",4,[
      ("Units index","148.3 (H1 = 100). Use consistent period and included transactions.","mint"),
      ("ASP index","66.9. Recalculate as net revenue / units with consistent weights.","amber"),
      ("Implied revenue index","148.3 × 66.9 / 100 = 99.2, or about -0.8%.","blue"),
      ("Reported revenue index","85.2, or -14.8%. The 14.0-point gap is a definition / weighting / calculation exception until proven otherwise.","red")])
    cards_slide(p,"04 / ANALYTICAL WORKPLAN","Driver hypotheses to test after reconciliation","These are hypotheses, not findings. Quantify each driver and distinguish correlation from causal evidence.",5,[
      ("Price & discount leakage","Test realized price vs list price, discount rate, promotion exposure and price bands. Evidence: line-item price, discount and units.","mint"),
      ("Product / customer mix","Check whether volume shifted toward lower-ticket products, bundles or segments. Evidence: SKU, category, basket and cohort.","blue"),
      ("Channel / regional mix","Compare unit economics across regions, store/digital channels, fulfillment and fee fields.","amber"),
      ("Returns, costs & metric logic","Validate returns, cancellations, cost allocation, timing, currency and gross/net definitions. Evidence: return ledger and metric lineage.","red")])
    cards_slide(p,"05 / EXECUTION PLAN","30-day decision plan","Sequence the work from conflicting KPIs to a defensible, testable management decision.",6,[
      ("WEEK 1 · RECONCILE","Metric dictionary; source-to-report trace; tie out revenue, units, ASP and profit.","mint"),
      ("WEEK 2 · DIAGNOSE","Price-volume-mix bridge; product, segment and channel cuts; anomaly review.","blue"),
      ("WEEK 3 · TEST","Design a controlled price / discount or mix intervention with explicit guardrails.","amber"),
      ("WEEK 4 · DECIDE","Executive readout comparing lift, downside, confidence and operating cost.","red")])
    cards_slide(p,"06 / REPORTING DESIGN","Executive dashboard: signal, quality, action","Leadership needs to know what changed, whether the numbers are trusted, and what decision comes next.",7,[
      ("OUTCOME","Net revenue · gross profit · contribution margin. Report value outcomes, not volume alone.","mint"),
      ("DRIVERS","Units · weighted ASP · discount rate. Drill down by product, region, segment and channel.","blue"),
      ("RISK / QUALITY","Returns rate · KPI reconciliation · data completeness. Flag definition mismatches before publication.","amber"),
      ("DECISION","Pilot lift · guardrail breaches · next action / owner. Name the reviewer and review date.","red")])
    cards_slide(p,"THE DECISION","Approve the reconciliation sprint.","Hold broad pricing changes until the KPI bridge is defensible.",8,[
      ("REQUEST FROM FINANCE / DATA","Transaction-level net revenue, units and selling price. Discounts, returns, cancellations, costs, periods, definitions and report filters.","mint"),
      ("APPROVAL TO MOVE FORWARD","Name Finance and Commercial owners. Agree metric definitions and success thresholds. Review a reconciled driver bridge before test launch.","blue")])
    p.save(OUT/"retail-sales-performance-diagnostic.pptx")
def pricing():
    p=deck("Pricing, Unit Economics & Scenario Planning")
    cover(p,"EXECUTIVE COMMERCIAL DECISION BRIEF","Pricing, Unit Economics\n& Scenario Planning","A quantified choice between growth, contribution, and profit protection","PRICE-UP CASE",
          [("Baseline operating profit","$100k","text"),("Price-up profit","$109k","mint"),"Assumed price +4%; unit volume -5%. Hypothetical inputs, not client results."])
    cards_slide(p,"01 / RECOMMENDATION","Executive takeaway","A price increase may improve profit despite lower revenue, but the demand response must be tested, not assumed.",2,[
      ("RECOMMENDED DECISION","Run a controlled price test before a broad rollout. Success metric: incremental operating contribution, with volume and retention guardrails.","mint"),
      ("REVENUE","$494k in the price-up case, down 1.2% from baseline. Revenue alone is not the success measure.","amber"),
      ("OPERATING PROFIT","$109k in the price-up case, up 9.0% under the stated assumptions. Contribution per unit rises from $20 to $22.","mint"),
      ("INVALIDATION / GUARDRAIL","Demand may be more price-sensitive; mix, returns or competitors may differ. Scale only if contribution improves and volume stays above break-even.","blue")])
    cards_slide(p,"02 / BASELINE","Baseline unit economics","Starting point for the scenario model. Replace every input with finance-approved actuals before operational use.",3,[
      ("PRICE / UNIT","$50 assumed selling price. Units per period: 10,000.","mint"),
      ("VARIABLE COST","$30 per unit assumed variable cost. Fixed cost: $100,000 per period.","blue"),
      ("PROFIT BRIDGE","Revenue $500,000 − variable costs $300,000 = contribution $200,000 − fixed costs $100,000 = operating profit $100,000.","mint"),
      ("UNIT CONTRIBUTION","$50 price − $30 variable cost = $20. Contribution margin = 40% of revenue.","amber")])
    cards_slide(p,"03 / SCENARIO COMPARISON","Compare the commercial choices","The discount-led growth case adds revenue but erodes contribution; the price-up case lifts profit under the stated demand assumption.",4,[
      ("BASE CASE","Price $50 · 10,000 units · $500,000 revenue · $200,000 contribution · $100,000 operating profit.","blue"),
      ("DISCOUNT + VOLUME","Price $48 · 11,000 units · $528,000 revenue · $198,000 contribution · $98,000 operating profit.","amber"),
      ("PRICE-UP TEST","Price $52 · 9,500 units · $494,000 revenue · $209,000 contribution · $109,000 operating profit.","mint")])
    cards_slide(p,"04 / BREAK-EVEN TEST","Break-even volume sets the guardrail","At $52 price and $30 variable cost, volume may decline about 9.1% and still preserve baseline operating profit.",5,[
      ("PRICE-UP CONTRIBUTION","$52 price − $30 variable cost = $22 contribution per unit.","mint"),
      ("VOLUME TO PRESERVE $100K PROFIT","($100k fixed costs + $100k target profit) / $22 = approximately 9,091 units.","blue"),
      ("BREAK-EVEN GUARDRAIL","About 90.9% of baseline volume is needed to preserve $100k operating profit under the assumptions.","amber"),
      ("ASSUMED CASE","95% volume retention = 9,500 units and $109k operating profit. Headroom is around 4.1 percentage points.","mint")])
    cards_slide(p,"05 / SENSITIVITY","Sensitivity: profit depends on retained units","At $52 price and $30 variable cost, profit = $22 × units − $100,000.",6,[
      ("85% VOLUME RETENTION","8,500 units · $442k revenue · $87k operating profit.","red"),
      ("90% VOLUME RETENTION","9,000 units · $468k revenue · $98k operating profit.","amber"),
      ("90.9% BREAK-EVEN","9,091 units · $473k revenue · $100k operating profit.","blue"),
      ("95% TO 105% RETENTION","$109k profit at 95%; $120k at 100%; $131k at 105%. Assumptions are not observed elasticities.","mint")])
    cards_slide(p,"06 / TEST PLAN","Pilot design: turn assumptions into evidence","Use a controlled test or justified quasi-experiment. Track contribution and customer outcomes, not revenue alone.",7,[
      ("01 · DESIGN","Choose comparable customer / product groups. Set treatment, control, duration, exclusions, and baselines.","mint"),
      ("02 · MEASURE","Track realized price, units, conversion, net revenue, contribution, returns, repeat purchase and complaints.","blue"),
      ("03 · GUARDRAIL","Set minimum contribution gain and maximum acceptable volume / retention decline; define stop rules.","amber"),
      ("04 · REVIEW","Estimate uplift with uncertainty; check segment effects, seasonality and operating costs before rollout.","red")])
    cards_slide(p,"DECISION MEMO","Test the price-up hypothesis.","Do not commit to a broad rollout yet. All dollar and demand-response assumptions are hypothetical.",8,[
      ("WHAT THE MODEL SAYS","Price $50 → $52 (+4%). Assumed units 10,000 → 9,500 (-5%). Revenue $500k → $494k (-1.2%). Operating profit $100k → $109k (+9.0%).","mint"),
      ("WHAT MANAGEMENT MUST PROVIDE","Actual price and cost data; demand / conversion history by segment; returns, retention and competitor context; pilot owner, review date and stop thresholds.","blue")])
    p.save(OUT/"pricing-unit-economics-scenario-planning.pptx")
if __name__=="__main__":
    retail(); pricing(); print(f"Built decks in {OUT}")
